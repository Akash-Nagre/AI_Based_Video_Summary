from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src = Path('AI_Video_Summarization_IEEE_Revised_v10.docx')
out = Path('AI_Video_Summarization_IEEE_Revised_v10_REVIEW_HIGHLIGHTED.docx')
doc = Document(src)

# Yellow marks the specific evidence and manuscript items that still need author action.
paragraph_starts = (
    'The project retains several source videos',
    'The selector comparison uses the same 909.8-s source',
    'All four outputs used the same 909.8-s Hindi source',
    'The supplied job records provide a same-source runtime comparison',
    '• Uniform, transcript-only, and vision-only selector outputs',
    '• Speech-recognition errors can propagate',
    'The next step is to collect source-aligned human reference intervals',
)
for p in doc.paragraphs:
    if p.text.startswith(paragraph_starts):
        for run in p.runs:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW

# Highlight pending status cells in the evaluation table.
status_table = doc.tables[3]
for row in status_table.rows[4:7]:
    for run in row.cells[2].paragraphs[0].runs:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW

# The current baseline table reports duration allocation only; mark its outcome rows for quality results.
for row in doc.tables[5].rows[1:]:
    for cell in row.cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

# Mark the run comparison caveat that needs either a controlled repeat or corresponding human evaluation.
for p in doc.paragraphs:
    if p.text.startswith('• Runtime measurements come from a small number of videos'):
        for run in p.runs:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW

# Add a document metadata note without inserting text into the paper body.
doc.core_properties.comments = 'Review copy: yellow highlighting marks evidence or author edits needed before submission.'
doc.save(out)
print(out.resolve())
