"""Revise the manuscript with the completed single-video selector comparison."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "AI_Video_Summarization_IEEE_Revised_v8.docx"
OUTPUT = ROOT / "AI_Video_Summarization_IEEE_Revised_v9.docx"


def replace_paragraph(paragraph, value: str):
    old_runs = paragraph.runs
    run_properties = deepcopy(old_runs[0]._r.rPr) if old_runs and old_runs[0]._r.rPr is not None else None
    for run in list(old_runs):
        paragraph._p.remove(run._r)
    new_run = paragraph.add_run(value)
    if run_properties is not None:
        new_run._r.insert(0, run_properties)


def set_cell(cell, value: str, *, bold=False, font_size=8):
    cell.text = value
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.bold = bold
            run.font.size = Pt(font_size)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def main():
    doc = Document(SOURCE)
    by_text = {p.text: p for p in doc.paragraphs}

    replacements = {
        "Abstract—Long recordings make it difficult to locate important passages, while a useful highlight should retain source timestamps and represent the main topics. This paper presents a local, extractive video-summarization pipeline that transcribes speech with Faster-Whisper, groups timestamped transcript segments into candidate clips, embeds their text, measures low-rate visual activity, and ranks clips using transcript importance, visual activity, and temporal position. A duration-budgeted greedy selector penalizes semantic redundancy and rewards uncovered transcript-derived topics and timeline regions; FFmpeg renders the selected clips. Content-fingerprint caching and bounded parallel rendering reduce repeat work. On a 909.8-s Hindi video, the default small-model, beam-3, batch-1 run produced a 227.34-s summary under a 227.45-s budget in 725.30 s (12m05s). Transcription took 579.33 s. Encoding the same selected intervals with two workers reduced clip-render time from 86.97 to 43.43 s; a same-source warm-cache run took 66.44 s and selected the same intervals. A batch-2 experiment completed in 105.17 s but produced substantially different transcripts, topics, and clips, so it was not adopted. These are runtime, duration, and internal-consistency measurements; no human reference summaries are available to establish semantic accuracy.":
            "Long recordings make it difficult to locate important passages, while useful highlights should preserve source timestamps and represent the main topics. This paper presents a local extractive video-summarization pipeline using Faster-Whisper, transcript embeddings, low-rate visual activity, and duration-budgeted topic-aware selection. Content-fingerprint caching and bounded parallel rendering reduce repeat work. On a 909.8-s Hindi video, the default small-model, beam-3, batch-1 run selected 227.34 s under a 227.45-s budget in 725.30 s. Two-worker clip rendering reduced the measured rendering stage from 86.97 to 43.43 s. We also generated fixed-budget uniform, transcript-only, and vision-only selector baselines on the same source. Their selected durations ranged from 227.36 to 227.45 s. Human reference intervals and reviewer ratings have not yet been evaluated, so these baseline runs do not establish comparative summary quality.",
        "The local evaluation corpus contains four unique videos (about 29 min 31 s in aggregate): English neuroscience and presentation/story videos and Hindi buoyancy and thermodynamics videos. Repeated job folders contain copies of these sources. The long-video runtime experiment used a 909.8-s Hindi thermodynamics lecture. Original source links were not retained. No human reference summaries, question sets, or answer-interval annotations are available, so semantic accuracy cannot be calculated from these runs.":
            "The project retains several source videos, including English and Hindi examples; repeated job folders contain copies of some sources. The selector-baseline comparison used the 909.8-s Hindi thermodynamics lecture and a 25% summary budget. Original source links were not retained. No source-aligned reference interval set or completed reviewer-rating data is included in this evaluation, so semantic accuracy cannot yet be calculated.",
        "The system was checked in four ways: (1) a 107-s English smoke test with a tiny Whisper model and a 15% summary budget; (2) stage-level profiling on the 909.8-s Hindi video with the default small-model, beam-3, batch-1 settings; (3) a same-source warm-cache rerun and a two-worker rendering comparison using the identical selected intervals; and (4) an exploratory batch-2 run compared with the default transcript and selection. Six focused automated regression tests passed, and FFmpeg decoded the generated long-video summary without error. These checks measure runtime and software behavior; they are not a human evaluation of summary quality.":
            "The system was checked in five ways: (1) a 107-s English smoke test with a tiny Whisper model and a 15% summary budget; (2) stage-level profiling on the 909.8-s Hindi video with the default small-model, beam-3, batch-1 settings; (3) a same-source warm-cache rerun and a two-worker rendering comparison using identical selected intervals; (4) an exploratory batch-2 run compared with the default transcript and selection; and (5) fixed-budget uniform, transcript-only, and vision-only selector baselines on the same 909.8-s source. Six focused automated regression tests passed. FFmpeg decoded the proposed and all three baseline outputs without error. These checks establish runtime, budget adherence, and output validity; human-judged summary quality remains unmeasured.",
        "The smoke test, runtime profiles, and cache/rendering comparisons verify software behavior, budget adherence, and repeatability for the tested inputs. They do not measure whether a viewer considers the selected content complete or useful.":
            "The smoke test, runtime profiles, and cache/rendering comparisons verify software behavior, budget adherence, and repeatability for the tested inputs. The new selector baselines make a same-source quality comparison possible, but no human quality scores have been collected.",
        "D.  Planned Evaluation":
            "D.  Selector Baseline and Human Evaluation Protocol",
        "A full evaluation would compare the generated summary intervals with human reference intervals using interval-level precision, recall and F1 at a temporal IoU of at least 0.5. An annotation tool for this metric has been implemented but has not yet been run against human annotations. Baselines should include uniform sampling and transcript-only and vision-only variants. Table IV lists the measures and their current status.":
            "The selector comparison uses the same 909.8-s source and 227.45-s duration budget for the proposed system and three baselines. Uniform sampling uses 21 equally spaced, equal-duration source windows. Transcript-only and vision-only use the same 57 candidate clips and a duration-constrained knapsack that maximizes the corresponding normalized score multiplied by clip duration; these variants omit the other modality and topic-coverage reward. For human evaluation, reference intervals will be matched one-to-one by maximum cardinality at temporal IoU ≥ 0.5, with interval precision, recall, and F1 reported. At least three independent reviewers will rate the four blinded outputs in randomized order for informativeness, topic coverage, coherence, factual consistency, and overall quality on a five-point scale; ordinal inter-rater agreement will also be reported. The interval and rating forms are prepared, but reference annotations and reviewer scores are not yet available; Table IV separates measured output checks from pending quality measures.",
        "Results include the 107-s software smoke test and the 909.8-s runtime profile on the hardware in Section IV-A. Table V reports the long-video profile, cache state, output checks, and the exploratory batch-2 comparison. All runs use one source video per setting; they are not a corpus-wide benchmark and do not measure human-judged summary quality.":
            "Results include the 107-s software smoke test and the 909.8-s long-video experiments on the hardware in Section IV-A. Table V reports runtime and optimization measurements; Table VI reports the proposed selector and three same-source baselines. The selector comparison covers one source video and reports interval allocation only. It is not a corpus-wide benchmark and does not measure human-judged summary quality.",
        "• No human-annotated evaluation, baseline comparison, or statistical analysis has been completed; the coverage score and transcript-disagreement checks do not establish summary accuracy.":
            "• Uniform, transcript-only, and vision-only selector outputs have been generated for one source video, but their human evaluation has not been completed. The heuristic coverage score does not establish summary accuracy.",
        "This paper described a local timestamped extractive video-summarization pipeline combining transcript importance, sampled visual activity, temporal position, and duration-budgeted topic-aware selection. On a 909.8-s Hindi source, the default small-model, beam-3, batch-1 run produced a 227.34-s output under the 227.45-s budget in 725.30 s; transcription accounted for 579.33 s. Reusing same-source analysis caches and encoding clips with two workers produced a 66.44-s repeat run, while the clip-render stage fell from 86.97 to 43.43 s with identical selected intervals and coverage. A faster batch-2 trial changed transcripts, topics, and selected intervals substantially, so it was not enabled. These results establish measured runtime and internal consistency only. Human-annotated evaluation and reference-based baselines remain necessary before making claims about semantic accuracy or viewer preference.":
            "This paper described a local timestamped extractive video-summarization pipeline combining transcript importance, sampled visual activity, temporal position, and duration-budgeted topic-aware selection. On a 909.8-s Hindi source, the default small-model, beam-3, batch-1 run selected 227.34 s under the 227.45-s budget in 725.30 s; transcription accounted for 579.33 s. The two-worker rendering stage fell from 86.97 to 43.43 s, and a same-source warm-cache repeat took 66.44 s. Uniform, transcript-only, and vision-only summaries were also generated within 0.09 s of the duration target. Human reference-interval metrics and reviewer ratings remain pending, so no claim of semantic superiority is made.",
        "The next step is an annotated evaluation: collecting human reference summaries for a set of videos, computing interval-level precision, recall and F1 at a temporal IoU of at least 0.5, and comparing the system with uniform-sampling, transcript-only and vision-only baselines, with ablations of the visual and positional scores and a sensitivity study of the fusion weights. Further work includes a visual-semantic encoder in place of frame differences, scene-based and speaker-aware segmentation, and, as an extension, timestamped retrieval and question answering [7]–[10], [18]. Fig. 5 shows the envisaged extension, which is not implemented in the present system.":
            "The next step is to collect source-aligned human reference intervals and independent reviewer ratings for the four generated summaries, then repeat the comparison across more videos and domains. Interval precision, recall, and F1 will be computed at temporal IoU ≥ 0.5; blinded ratings will cover informativeness, topic coverage, coherence, factual consistency, and overall quality. Future work also includes ablations of the visual and positional signals, a sensitivity study of selector weights, a visual-semantic encoder, scene-aware and speaker-aware segmentation, and, as an extension, timestamped retrieval and question answering [7]–[10], [18]. Fig. 5 shows the envisaged extension, which is not implemented in the present system.",
    }
    for old, new in replacements.items():
        if old not in by_text:
            raise ValueError(f"Expected manuscript paragraph not found: {old[:90]}")
        replace_paragraph(by_text[old], new)

    # Update status table and include subjective human ratings as a separate measure.
    evaluation_table = doc.tables[3]
    table_rows = {row.cells[0].text: row for row in evaluation_table.rows}
    table_rows["Interval F1"].cells[2].text = "Annotation template prepared; source-aligned intervals and scores pending"
    table_rows["Interval F1"].cells[1].text = "One-to-one precision, recall, F1 at temporal IoU ≥ 0.5"
    table_rows["Baselines and ablations"].cells[2].text = "Three selector baselines generated on the same 909.8-s source; quality scoring pending"
    human_row = evaluation_table.add_row()
    human_row.cells[0].text = "Human ratings"
    human_row.cells[1].text = "Five-point informativeness, coverage, coherence, factual consistency, and overall quality"
    human_row.cells[2].text = "Blind review packet prepared; reviewer scores pending"

    # Add a compact, single-column table before Section VI.
    reference = next(p for p in doc.paragraphs if p.text == "VI.  LIMITATIONS")
    heading_template = next(p for p in doc.paragraphs if p.text.startswith("D. Faster-Inference Trial"))
    body_template = next(p for p in doc.paragraphs if p.text.startswith("An exploratory batch-2 run"))
    heading = doc.add_paragraph(style=heading_template.style)
    if heading_template._p.pPr is not None:
        heading._p.insert(0, deepcopy(heading_template._p.pPr))
    hrun = heading.add_run("E. Selector Baseline Outputs")
    if heading_template.runs and heading_template.runs[0]._r.rPr is not None:
        hrun._r.insert(0, deepcopy(heading_template.runs[0]._r.rPr))

    intro = doc.add_paragraph(style=body_template.style)
    if body_template._p.pPr is not None:
        intro._p.insert(0, deepcopy(body_template._p.pPr))
    intro.add_run(
        "All four outputs used the same 909.8-s Hindi source and 227.45-s target budget. The proposed selector used its saved batch-1 candidate scores and topic-coverage rewards. Uniform used equal windows spaced across the full source; transcript-only and vision-only ranked the common set of 57 candidates by one signal using a 10-ms duration-knapsack. Table VI reports selected source-interval duration, not the encoded MP4 container duration. The selected intervals met the budget for every method. These measurements verify comparable duration allocation, not semantic quality; human interval and rating scores remain pending."
    )

    baseline_table = doc.add_table(rows=5, cols=3)
    baseline_table.style = doc.tables[3].style
    baseline_table.autofit = False
    widths = [Inches(0.97), Inches(1.41), Inches(0.87)]
    values = [
        ("Method", "Selection rule", "Selected source duration and clips"),
        ("Proposed", "Fused salience, diversity, and topic rewards", "227.34 s; 21"),
        ("Uniform", "Evenly spaced equal-duration windows", "227.45 s; 21"),
        ("Transcript-only", "Transcript score × duration knapsack", "227.42 s; 18"),
        ("Vision-only", "Visual score × duration knapsack", "227.36 s; 20"),
    ]
    for row_idx, data in enumerate(values):
        row = baseline_table.rows[row_idx]
        prevent_row_split(row)
        for col_idx, value in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = widths[col_idx]
            set_cell(cell, value, bold=(row_idx == 0), font_size=7.5)

    # Apply the manuscript's table caption convention before the table.
    caption_template = next(p for p in doc.paragraphs if p.text.startswith("TABLE V."))
    caption = doc.add_paragraph(style=caption_template.style)
    if caption_template._p.pPr is not None:
        caption._p.insert(0, deepcopy(caption_template._p.pPr))
    caption.add_run("TABLE VI. SAME-SOURCE SELECTOR BASELINE OUTPUTS")

    for element in (heading._p, intro._p, caption._p, baseline_table._tbl):
        reference._p.addprevious(element)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
